"""
Freyja pre-processor
====================
Builds models/freyja.xml (the MuJoCo model) from models/freyja_template.xml by
pulling numbers out of the Freyja Anthropometric Reference Google Sheet.

    ground-truth sheet  -->  this script  -->  freyja.xml
    (named ranges)           (check + fill)     (never edit by hand!)

SETUP (once): put the sheet ID and the path to the service-account key in
    local_config.json (copy local_config.example.json). That file is git-ignored.
    Or set the environment variables FREYJA_SHEET_KEY and FREYJA_CREDENTIALS.

TO REBUILD AUTOMATICALLY whenever the sheet changes, leave scripts\\watch.py running
    (it imports build_model() from this file, so both always behave the same).

HOW TO RUN (from the sim/ folder, using the project's virtual environment)
    freyja.venv\\Scripts\\python scripts\\pre_processor.py
    freyja.venv\\Scripts\\python scripts\\pre_processor.py --xlsx "C:\\path\\to\\sheet.xlsx"
        works offline from a downloaded copy of the sheet (needs: pip install openpyxl)
    freyja.venv\\Scripts\\python scripts\\pre_processor.py --strict    fail if any ROM is missing
    freyja.venv\\Scripts\\python scripts\\pre_processor.py --dry-run   check everything, write nothing

WHAT IT PULLS (three named ranges in the sheet: Data > Named ranges)
    bsip : mass, length, CoM, inertia tensor per segment
    pos  : body positions (pos_x/y/z) and joint anchors (jpos_z) per segment
    rom  : joint range-of-motion limits (the "MuJoCo limit - FINAL" column)
    Segments are looked up by NAME and columns by HEADER, never by row/column
    number, so inserting rows or reordering columns in the sheet is safe.
    Renaming a segment or a header is NOT safe: the script will say so.

HOW TO WRITE PLACEHOLDERS IN THE TEMPLATE
    ${thigh_right_mass}  or  !{thigh_right_mass}      (both styles work)
        BSIP/position names are  <segment>_<value>, e.g.
        pelvis_com_x   thigh_left_ixy   upper_arm_right_pos_z   abdomen_jpos_z
        Bilateral segments (upper_arm forearm hand thigh shank foot) exist as
        _right (straight from the sheet) and _left (mirrored in y automatically).
    !{hip_flexion}   ROM names are  <joint>_<movement>  exactly as in the sheet,
        lower-case with underscores: hip_internal_rotation, ankle_dorsiflexion ...
    !{-hip_extension}   a leading minus negates the value. ROM values are stored
        as positive magnitudes, so the template decides the sign. Because each
        joint axis already fixes which movement is positive, the sign convention
        sits right next to the axis it belongs to:
            axis="0 -1 0" range="!{-hip_extension} !{hip_flexion}"

WHEN A VALUE IS MISSING
    BSIP / position values:  always an error (there is no sane default).
    ROM values:             the joint is left unlimited, a FIXME comment is put in
                            the XML and a warning is printed (same thing you did
                            by hand for the spine). Use --strict to make it fatal.

EVERY RUN ALSO
    * checks the numbers (formats, positive mass, valid inertia tensors, chain
      lengths, ROM sanity) and reports ALL problems at once with the sheet cell
    * checks the finished XML is well-formed and that MuJoCo can compile it
    * never replaces a good freyja.xml with a broken one
    * saves a CSV of every value used to archive/ (only when something changed)
    * installs params/snapshot.csv and snapshot.meta.json together with the model
    * runs the model checks (checks/): a failing gate check parks the new XML as freyja_FAILED.xml and
      keeps the last good model and snapshot; advisory results are printed and saved to checks/last_run.json
"""

from __future__ import annotations

import argparse
import csv
import difflib
import hashlib
import io
import json
import math
import os
import re
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np

# =============================================================================
# SETTINGS  (the only section you should normally need to touch)
# =============================================================================
HERE = Path(__file__).resolve().parent
ROOT = HERE.parent  # the sim/ folder

# The sheet ID and the service-account key path are NOT stored in this file.
# They come from (first match wins):
#   1. environment variables  FREYJA_SHEET_KEY and FREYJA_CREDENTIALS
#   2. local_config.json in the sim/ folder (git-ignored; copy
#      local_config.example.json to start one)
CONFIG_FILE = ROOT / "local_config.json"
SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]  # read-only

TEMPLATE = ROOT / "models" / "freyja_template.xml"
OUTPUT = ROOT / "models" / "freyja.xml"
FAILED_OUTPUT = ROOT / "models" / "freyja_FAILED.xml"  # a broken build is parked here
ARCHIVE_DIR = ROOT / "archive"
PARAMS_DIR = ROOT.parent / "params"  # snapshot.csv + snapshot.meta.json (CONVENTIONS section 5)
SNAPSHOT_SCHEMA = "params/snapshot/1"

# Names of the named ranges in the Google Sheet
RANGE_BSIP, RANGE_POS, RANGE_ROM = "bsip", "pos", "rom"
# Optional: total mass and stature (label in column A, value in B, unit in the
# label's brackets). Feeds the snapshot only; the template cannot use them.
RANGE_TARGETS = "targets"
TARGET_KEYS = {"mass": "target_mass", "stature": "target_stature"}  # label keyword -> snapshot key
TARGET_UNIT_TO_SI = {"kg": 1.0, "g": 1e-3, "m": 1.0, "mm": 1e-3, "cm": 1e-2}  # the snapshot is SI

# The sheet must contain exactly these segments (after ident()). If someone adds,
# renames or deletes one in the sheet the script stops and says so ("sheet drift").
EXPECTED_SEGMENTS = ["head_neck", "thorax", "abdomen", "pelvis", "upper_arm",
                     "forearm", "hand", "thigh", "shank", "foot"]
BILATERAL = ["upper_arm", "forearm", "hand", "thigh", "shank", "foot"]

# Column headers (units in brackets are ignored: "mass (kg)" -> "mass")
BSIP_COLUMNS = ["mass", "length", "com_x", "com_y", "com_z",
                "ixx", "iyy", "izz", "ixy", "ixz", "iyz"]
POS_COLUMNS = ["pos_x", "pos_y", "pos_z", "jpos_z"]

# Mirroring the right side into the left (MuJoCo +y = left). Only values that
# involve y change sign: CoM y, the products of inertia containing y, body pos y.
FLIP_BSIP = {"com_y", "ixy", "iyz"}
FLIP_POS = {"pos_y"}

# In the ROM table, which column holds the value MuJoCo should use
ROM_VALUE_HEADER_KEYWORD = "final"  # matches "MuJoCo limit - FINAL (deg)"
ROM_MAX_DEG = 360.0

# Each body sits one parent-segment-length from its parent: (child, segment whose
# length it should equal). A mismatch only warns, since you may do it on purpose.
CHAIN_CHECKS = [("abdomen", "abdomen"), ("thorax", "thorax"),
                ("forearm_right", "upper_arm_right"), ("hand_right", "forearm_right"),
                ("shank_right", "thigh_right"), ("foot_right", "shank_right")]
CHAIN_TOLERANCE_M = 0.001

# =============================================================================
# SMALL HELPERS
# =============================================================================


def ident(name) -> str:
    """'Head & Neck' -> 'head_neck'. Turns any sheet label into a safe key."""
    return re.sub(r"\W+", "_", str(name).strip().lower()).strip("_")


def column_key(header) -> str:
    """'CoM_x (m)' -> 'com_x', 'mass (kg)' -> 'mass'. Bracketed units are dropped."""
    if header is None:
        return ""
    return ident(re.sub(r"\(.*?\)", "", str(header)))


def fmt(v) -> str:
    """Number -> text for the XML. Also removes the -0.0 ambiguity."""
    x = float(v)
    if x == 0.0:
        x = 0.0
    return f"{x:.12g}"


def to_number(value):
    """Cell -> float, or None if the cell is blank. Raises ValueError if it is
    text, an error like #REF!, NaN or infinity."""
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    if isinstance(value, bool):
        raise ValueError(value)
    x = float(value)  # ValueError for text
    if not math.isfinite(x):
        raise ValueError(value)
    return x


def col_letter(n: int) -> str:
    """1 -> A, 27 -> AA"""
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


class BuildError(Exception):
    """The model could not be built. Carries every problem found, so callers
    (the command line, the watcher) decide how to show it. Never kills the process."""

    def __init__(self, stage: str, errors: list, warnings: list, transient: bool = False):
        super().__init__(f"{len(errors)} problem(s) found while {stage}")
        self.stage, self.errors, self.warnings = stage, list(errors), list(warnings)
        self.transient = transient  # True = the data was fine, the environment wasn't (e.g. file locked)

    def format(self) -> str:
        lines = [f"Stopped: {len(self.errors)} problem(s) found while {self.stage}"]
        lines += [f"  ERROR   {e}" for e in self.errors]
        lines += [f"  WARNING {w}" for w in self.warnings]
        return "\n".join(lines)


class SheetAccessError(Exception):
    """Could not read the sheet (network, permissions, bad name ...). `fatal`
    means retrying cannot help (e.g. no credentials configured)."""

    def __init__(self, message: str, fatal: bool = False):
        super().__init__(message)
        self.fatal = fatal


class Report:
    """Collects every problem so you see them all at once, not one per run."""

    def __init__(self):
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, msg: str):
        self.errors.append(msg)

    def warn(self, msg: str):
        self.warnings.append(msg)

    def stop_if_errors(self, stage: str):
        if self.errors:
            raise BuildError(stage, self.errors, self.warnings)


# =============================================================================
# 1. GETTING THE RAW TABLES (Google Sheet, or a downloaded .xlsx for offline use)
# =============================================================================


@dataclass
class RawTable:
    """One named range exactly as it came from the sheet (header row first)."""
    name: str
    sheet: str
    first_row: int  # sheet row number of rows[0]
    first_col: int  # sheet column number of rows[0][0]
    rows: list

    @classmethod
    def from_range(cls, name, range_str, rows):
        m = re.match(r"^'?(.+?)'?!\$?([A-Z]+)\$?(\d+)", range_str)
        if not m:
            raise ValueError(f"cannot understand range '{range_str}'")
        col = 0
        for ch in m.group(2):
            col = col * 26 + ord(ch) - 64
        return cls(name, m.group(1), int(m.group(3)), col, rows)

    def ref(self, row_idx: int, col_idx: int) -> str:
        """Where a cell lives in the sheet, e.g. 'MuJoCo Reference'!G20"""
        sheet = f"'{self.sheet}'" if " " in self.sheet else self.sheet
        return f"{sheet}!{col_letter(self.first_col + col_idx)}{self.first_row + row_idx}"


def setting(env_name: str, config_key: str):
    """Look a private setting up in the environment, then in local_config.json."""
    value = os.environ.get(env_name)
    if value:
        return value
    if CONFIG_FILE.exists():
        try:
            cfg = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
        except (OSError, ValueError) as e:
            raise SheetAccessError(f"{CONFIG_FILE.name} could not be read ({e}). "
                                   "It must be valid JSON, see local_config.example.json.", fatal=True)
        return cfg.get(config_key) or None
    return None


class GoogleSheetSource:
    """Logs in once, then each fetch() is a single small request for the three
    named ranges, so it is cheap enough to poll every few seconds."""
    REQUEST_TIMEOUT_S = 30  # a hung connection must not freeze the watcher

    def __init__(self):
        self.sheet_key = setting("FREYJA_SHEET_KEY", "sheet_key")
        cred_path = setting("FREYJA_CREDENTIALS", "credentials")
        if not self.sheet_key or not cred_path:
            raise SheetAccessError(
                "The sheet ID and/or credentials path are not configured.\n"
                f"  Copy local_config.example.json to {CONFIG_FILE.name} (in the sim/ folder) and fill it in,\n"
                "  or set the environment variables FREYJA_SHEET_KEY and FREYJA_CREDENTIALS.", fatal=True)
        if not Path(cred_path).exists():
            raise SheetAccessError(f"Credentials file not found:\n  {cred_path}\n"
                                   "Fix 'credentials' in local_config.json (or FREYJA_CREDENTIALS).", fatal=True)
        import gspread  # imported here so --xlsx and the tests work without it
        from google.oauth2.service_account import Credentials
        from gspread.http_client import BackOffHTTPClient  # waits and retries on rate limits (429)
        self._gspread = gspread
        self._creds = Credentials.from_service_account_file(cred_path, scopes=SCOPES)
        self._client = gspread.authorize(self._creds, http_client=BackOffHTTPClient)
        self._client.set_timeout(self.REQUEST_TIMEOUT_S)
        self._sheet = None

    def fetch(self) -> dict:
        gspread = self._gspread
        email = self._creds.service_account_email
        try:
            if self._sheet is None:
                self._sheet = self._client.open_by_key(self.sheet_key)
            params = {"valueRenderOption": "UNFORMATTED_VALUE"}
            names = [RANGE_BSIP, RANGE_POS, RANGE_ROM]
            try:
                resp = self._sheet.values_batch_get(names + [RANGE_TARGETS], params=params)
                names.append(RANGE_TARGETS)
            except gspread.exceptions.APIError:  # no 'targets' range yet: it is optional
                resp = self._sheet.values_batch_get(names, params=params)
        except gspread.exceptions.SpreadsheetNotFound:
            raise SheetAccessError("Google says the sheet doesn't exist or isn't shared with this account. "
                                   f"Share it (Viewer is enough) with: {email}")
        except gspread.exceptions.APIError as e:
            status = getattr(getattr(e, "response", None), "status_code", "?")
            hint = {400: f"One of the named ranges ({RANGE_BSIP}, {RANGE_POS}, {RANGE_ROM}) "
                         "was not found. Check Data > Named ranges in the sheet.",
                    403: f"Not allowed. Share the sheet with: {email}",
                    404: "Sheet not found. Check the sheet ID in local_config.json."}.get(status, "")
            raise SheetAccessError(f"Google Sheets API error {status}. {hint} {e}".replace("  ", " "))
        except Exception as e:  # network down, DNS, timeouts ...
            raise SheetAccessError(f"Could not reach Google Sheets ({type(e).__name__}: {e})")

        tables = {}
        for name, vr in zip(names, resp["valueRanges"]):
            tables[name] = RawTable.from_range(name, vr["range"], vr.get("values", []))
        return tables


def fetch_google() -> dict:
    return GoogleSheetSource().fetch()


def fetch_xlsx(path: str) -> dict:
    try:
        import openpyxl
    except ImportError:
        raise SheetAccessError("Reading an .xlsx needs openpyxl:  freyja.venv\\Scripts\\pip install openpyxl", fatal=True)
    if not Path(path).exists():
        raise SheetAccessError(f"Spreadsheet file not found: {path}", fatal=True)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)  # data_only = the computed values
    except Exception as e:  # e.g. Excel is in the middle of saving the file
        raise SheetAccessError(f"could not read {path} ({type(e).__name__}: {e})")
    tables = {}
    for name in (RANGE_BSIP, RANGE_POS, RANGE_ROM):
        if name not in wb.defined_names:
            raise SheetAccessError(f"Named range '{name}' not found in {path}. "
                                   f"Found: {list(wb.defined_names.keys())}")
        dest = wb.defined_names[name].attr_text  # 'MuJoCo Reference'!$A$14:$L$24
        sheet, rng = dest.rsplit("!", 1)
        ws = wb[sheet.strip("'")]
        rows = [["" if c.value is None else c.value for c in row]
                for row in ws[rng.replace("$", "")]]
        tables[name] = RawTable.from_range(name, f"{sheet}!{rng}", rows)
    if RANGE_TARGETS in wb.defined_names:  # optional
        sheet, rng = wb.defined_names[RANGE_TARGETS].attr_text.rsplit("!", 1)
        ws = wb[sheet.strip("'")]
        rows = [["" if c.value is None else c.value for c in row] for row in ws[rng.replace("$", "")]]
        tables[RANGE_TARGETS] = RawTable.from_range(RANGE_TARGETS, f"{sheet}!{rng}", rows)
    return tables


# =============================================================================
# 2. TURNING TABLES INTO NUMBERS (and checking them)
# =============================================================================


@dataclass
class Param:
    value: float
    source: str  # where it came from, for error messages and the CSV archive
    table: str = ""  # bsip / pos / rom


def read_segment_table(raw: RawTable, columns: list, report: Report) -> dict:
    """bsip / pos table -> {segment: {column: Param}}. Checks segment names,
    headers and that every needed cell is a real number."""
    if not raw.rows:
        report.error(f"[{raw.name}] the named range is empty")
        return {}
    header = [column_key(h) for h in raw.rows[0]]
    if "segment" not in header:
        report.error(f"[{raw.name}] no 'Segment' header found. Headers seen: {raw.rows[0]}")
        return {}
    seg_i = header.index("segment")

    col_i = {}
    for col in columns:
        if col in header:
            col_i[col] = header.index(col)
        else:
            report.error(f"[{raw.name}] column '{col}' not found. "
                         f"Headers seen: {[h for h in header if h]}")

    found = {}
    for r, row in enumerate(raw.rows[1:], start=1):
        label = row[seg_i] if seg_i < len(row) else ""
        if str(label).strip() == "":
            continue
        seg = ident(label)
        if seg in found:
            report.error(f"[{raw.name}] segment '{label}' appears twice "
                         f"(second one at {raw.ref(r, seg_i)})")
            continue
        params = {}
        for col, ci in col_i.items():
            cell = row[ci] if ci < len(row) else ""
            ref = raw.ref(r, ci)
            try:
                x = to_number(cell)
            except (ValueError, TypeError):
                report.error(f"[{raw.name}] {ref} ({label} / {col}): {cell!r} is not a number")
                continue
            if x is None:
                report.error(f"[{raw.name}] {ref} ({label} / {col}): cell is empty")
                continue
            params[col] = Param(x, ref, raw.name)
        found[seg] = params

    missing = [s for s in EXPECTED_SEGMENTS if s not in found]
    extra = [s for s in found if s not in EXPECTED_SEGMENTS]
    if missing or extra:
        report.error(f"[{raw.name}] sheet drift. Missing segments: {missing or 'none'}; "
                     f"unexpected segments: {extra or 'none'}. Recent sheet edit vs "
                     "EXPECTED_SEGMENTS in pre_processor.py?")
    return found


def read_rom_table(raw: RawTable, report: Report):
    """rom table -> ({'hip_flexion': 133.8, ...}, {names that exist but are blank})"""
    values: dict[str, Param] = {}
    blank: set[str] = set()
    if not raw.rows:
        report.error(f"[{raw.name}] the named range is empty")
        return values, blank
    header = [column_key(h) for h in raw.rows[0]]
    for needed in ("joint", "movement"):
        if needed not in header:
            report.error(f"[{raw.name}] no '{needed}' header found. Headers seen: {raw.rows[0]}")
            return values, blank
    final = [i for i, h in enumerate(header) if ROM_VALUE_HEADER_KEYWORD in h]
    if len(final) != 1:
        report.error(f"[{raw.name}] expected exactly one column whose header contains "
                     f"'{ROM_VALUE_HEADER_KEYWORD}', found {len(final)}. Headers: {raw.rows[0]}")
        return values, blank
    ji, mi, vi = header.index("joint"), header.index("movement"), final[0]

    functional_i = header.index("functional") if "functional" in header else None
    for r, row in enumerate(raw.rows[1:], start=1):
        cells = lambda i: row[i] if i < len(row) else ""  # noqa: E731
        if str(cells(ji)).strip() == "" or str(cells(mi)).strip() == "":
            continue
        key = ident(f"{cells(ji)} {cells(mi)}")
        if key in values or key in blank:
            report.error(f"[{raw.name}] '{cells(ji)} / {cells(mi)}' appears twice "
                         f"(second one at {raw.ref(r, mi)})")
            continue
        ref = raw.ref(r, vi)
        try:
            x = to_number(cells(vi))
        except (ValueError, TypeError):
            report.error(f"[{raw.name}] {ref} ({key}): {cells(vi)!r} is not a number")
            continue
        if x is None:
            blank.add(key)
            continue
        if not 0 <= x <= ROM_MAX_DEG:
            report.error(f"[{raw.name}] {ref} ({key}): {x:g} deg is outside 0..{ROM_MAX_DEG:g}. "
                         "ROM values are positive magnitudes; the template supplies the sign.")
            continue
        values[key] = Param(x, ref, raw.name)
        if functional_i is not None:  # soft check: can the limit even do the task?
            try:
                demand = to_number(cells(functional_i))
            except (ValueError, TypeError):
                demand = None  # the column also holds notes like 'low' or '~0'
            if demand is not None and demand > x:
                report.warn(f"[{raw.name}] {key}: MuJoCo limit {x:g} deg is below the "
                            f"functional demand {demand:g} deg")
    return values, blank


def read_targets(raw, report: Report) -> dict:
    """targets table (label | value) -> {'target_mass': Param (kg), 'target_stature': Param (m)}.
    Optional and advisory: a problem here is a warning and the target is left
    out, so later checks that need it skip."""
    out: dict[str, Param] = {}
    if raw is None:
        return out
    for r, row in enumerate(raw.rows):
        label = str(row[0]) if row else ""
        key = next((k for word, k in TARGET_KEYS.items() if word in ident(label)), None)
        if key is None:
            if label.strip():
                report.warn(f"[{raw.name}] {raw.ref(r, 0)}: label '{label}' is not a known target; ignored")
            continue
        ref = raw.ref(r, 1)
        unit = re.search(r"\(\s*([A-Za-z]+)\s*\)", label)
        scale = TARGET_UNIT_TO_SI.get(unit.group(1).lower()) if unit else None
        try:
            x = to_number(row[1] if len(row) > 1 else None)
        except (ValueError, TypeError):
            x = None
        if scale is None or x is None or x <= 0:
            report.warn(f"[{raw.name}] {ref} ({label}): needs a positive number and a unit in brackets "
                        f"({', '.join(TARGET_UNIT_TO_SI)}); {key} left out of the snapshot")
            continue
        out[key] = Param(x * scale, ref, raw.name)
    return out


def check_bsip(bsip: dict, report: Report):
    """Physical sanity of each segment's mass properties (before mirroring)."""
    for seg, p in bsip.items():
        if len(p) < len(BSIP_COLUMNS):
            continue  # a missing/bad cell was already reported
        v = {k: x.value for k, x in p.items()}
        where = f"[bsip] {seg} (row source {p['mass'].source})"
        if v["mass"] <= 0:
            report.error(f"{where}: mass must be > 0, got {v['mass']:g}")
        if v["length"] <= 0:
            report.error(f"{where}: length must be > 0, got {v['length']:g}")
        I = np.array([[v["ixx"], v["ixy"], v["ixz"]],
                      [v["ixy"], v["iyy"], v["iyz"]],
                      [v["ixz"], v["iyz"], v["izz"]]])
        a, b, c = np.linalg.eigvalsh(I)  # principal moments, ascending
        if a <= 0:
            report.error(f"{where}: inertia tensor is not positive definite "
                         f"(smallest principal moment {a:.3g}). Check the signs of the products of inertia.")
        elif a + b < c * (1 - 1e-9):
            report.error(f"{where}: principal moments {a:.4g}, {b:.4g}, {c:.4g} break the triangle "
                         "inequality (a + b >= c). No real object has this inertia and MuJoCo will refuse it.")
        com = math.sqrt(v["com_x"] ** 2 + v["com_y"] ** 2 + v["com_z"] ** 2)
        if v["length"] > 0 and com > v["length"]:
            report.warn(f"{where}: centre of mass is {com:.3f} m from the origin, "
                        f"further than the segment is long ({v['length']:.3f} m)")


def mirror(params: dict, flip: set) -> dict:
    """Bilateral segments become <seg>_right (as in the sheet) and <seg>_left
    (mirrored in y). Everything else is passed through unchanged."""
    out = {}
    for seg, p in params.items():
        if seg in BILATERAL:
            out[f"{seg}_right"] = dict(p)
            out[f"{seg}_left"] = {k: Param(-q.value if k in flip else q.value,
                                           q.source + " (mirrored)", q.table)
                                  for k, q in p.items()}
        else:
            out[seg] = p
    return out


def flatten(params: dict) -> dict:
    """{'thigh_right': {'mass': Param}} -> {'thigh_right_mass': Param}"""
    return {f"{seg}_{key}": q for seg, p in params.items() for key, q in p.items()}


def check_positions(values: dict, report: Report) -> float | None:
    """Checks on the (mirrored) numbers. Returns the standing ankle height (m)."""
    v = {k: p.value for k, p in values.items()}
    for k, x in v.items():
        if "pos" in k and abs(x) > 2.0:
            report.error(f"[pos] {k} = {x:g} m looks wrong (units? expected metres)")

    for child, parent in CHAIN_CHECKS:
        try:
            gap = abs(v[f"{child}_pos_z"]) - v[f"{parent}_length"]
        except KeyError:
            continue
        if abs(gap) > CHAIN_TOLERANCE_M:
            report.warn(f"[pos] {child} sits {abs(v[f'{child}_pos_z']):.4f} m from its parent but "
                        f"{parent} is {v[f'{parent}_length']:.4f} m long (off by {gap * 1000:+.1f} mm)")

    for k in ("thigh_right_pos_y", "upper_arm_right_pos_y"):
        if k in v and v[k] >= 0:
            report.warn(f"[pos] {k} = {v[k]:g}: right-side bodies should have NEGATIVE y "
                        "(MuJoCo +y is the subject's left)")

    try:
        ankle = sum(v[f"{s}_pos_z"] for s in ("pelvis", "thigh_right", "shank_right", "foot_right"))
    except KeyError:
        return None
    if ankle < 0:
        report.error(f"[pos] standing chain puts the ankle {ankle * 1000:.1f} mm BELOW the floor "
                     "(pelvis + thigh + shank + foot heights should sum to the ankle height)")
    elif ankle > 0.15:
        report.warn(f"[pos] ankle is {ankle * 1000:.0f} mm above the floor in the standing pose; expected ~60-70 mm")
    return ankle


# =============================================================================
# 3. FILLING IN THE TEMPLATE
# =============================================================================

# ${name} and !{name} are the same thing; a leading '-' inside negates the value.
ANY_BRACES = re.compile(r"([$!])\{([^}]*)\}")
VALID_NAME = re.compile(r"^\s*(-?)\s*([A-Za-z_][A-Za-z0-9_]*)\s*$")
JOINT_TAG = re.compile(r"<joint\b[^<>]*>")
RANGE_ATTR = re.compile(r'\s+range="([^"]*)"')


@dataclass
class Placeholder:
    start: int
    end: int
    line: int
    negate: bool
    key: str


def find_placeholders(text: str, report: Report | None = None) -> list:
    found = []
    for m in ANY_BRACES.finditer(text):
        line = text.count("\n", 0, m.start()) + 1
        ok = VALID_NAME.match(m.group(2))
        if not ok:
            if report is not None:
                report.error(f"template line {line}: placeholder '{m.group(0)}' has no valid name. "
                             "Put a value name inside, e.g. !{hip_flexion}")
            continue
        found.append(Placeholder(m.start(), m.end(), line, ok.group(1) == "-", ok.group(2)))
    return found


def unlimit_joints(text: str, missing: set, report: Report) -> str:
    """For every <joint> whose range= uses a value that is missing: drop
    limited/range and leave a FIXME comment (what you did by hand for the spine)."""

    def fix(m):
        tag = m.group(0)
        rng = RANGE_ATTR.search(tag)
        if not rng:
            return tag
        gone = sorted({p.key for p in find_placeholders(rng.group(1)) if p.key in missing})
        if not gone:
            return tag
        name = re.search(r'name="([^"]*)"', tag)
        name = name.group(1) if name else "?"
        tag = re.sub(r'\s+limited="true"', "", RANGE_ATTR.sub("", tag))
        report.warn(f"ROM missing for joint '{name}' (template line "
                    f"{text.count(chr(10), 0, m.start()) + 1}): {', '.join(gone)}. "
                    "Joint left UNLIMITED until the sheet has the value.")
        return f"{tag} <!-- FIXME(pre_processor): no ROM in sheet for {', '.join(gone)}; joint left unlimited -->"

    return JOINT_TAG.sub(fix, text)


def render(text: str, values: dict, report: Report, strict: bool = False) -> str:
    return render_counted(text, values, report, strict)[0]


def render_counted(text: str, values: dict, report: Report, strict: bool = False):
    """Replace every placeholder. Unknown names are collected into the report.
    Returns (finished text, how many placeholders were filled in)."""
    phs = find_placeholders(text, report)
    missing = {p.key for p in phs if p.key not in values}
    if missing and not strict:
        text = unlimit_joints(text, missing, report)
        phs = find_placeholders(text, report)
        missing = {p.key for p in phs if p.key not in values}

    for p in phs:
        if p.key in missing:
            hint = difflib.get_close_matches(p.key, values.keys(), n=3, cutoff=0.7)
            report.error(f"template line {p.line}: '{p.key}' is not in the sheet data"
                         + (f". Did you mean: {', '.join(hint)}?" if hint else ""))
    if report.errors:
        return text, 0

    def sub(m):
        ok = VALID_NAME.match(m.group(2))
        x = values[ok.group(2)]
        return fmt(-x if ok.group(1) == "-" else x)

    return ANY_BRACES.subn(sub, text)


# Derived keys: sheet values combined in the one way the template geometry needs.
# The placeholder language stays closed; these are the only names that are not in the sheet.
# A capsule's end spheres stick out by one radius, so the centres sit one radius inside the
# joints; that is why geometry needs the radii here. All radii are design choices.
CAPSULE_RADII_M = {"thigh": 0.05, "shank": 0.04, "upper_arm": 0.03, "forearm": 0.02}  # design choice, no sheet source
KNEE_SPHERE_RADIUS_M = 0.045  # design choice, no sheet source: sphere centred on the knee
PELVIS_RADIUS_M = 0.055  # design choice, no sheet source
NECK_RADIUS_M = 0.045  # design choice, no sheet source
FOOT_REAR_FRACTION = 0.25  # design choice, no sheet source: the box reaches this fraction of the foot length behind the ankle
HEAD_HALF_HEIGHT_M = 0.095  # design choice, no sheet source: the head is an ellipsoid, top at the vertex
SHOULDER_RADIUS_M = 0.045  # design choice, no sheet source: shoulder girdle capsule between the shoulder joints
TRAPEZIUS_RADIUS_M = 0.04  # design choice, no sheet source
CHEST_Z_FRACTION = 0.30  # design choice, no sheet source: upper-chest ellipsoid centre, as a fraction of thorax length below the top
CHEST_HALF_FRACTION = 0.20  # design choice, no sheet source: its half height, same units
# Body origin at the distal end, segment hangs below. The value is the ellipsoid half length as a
# fraction of the segment length: 0.5 ends exactly at the joints; more overlaps the neighbours so
# the silhouette has no pinch at the joint (all these geoms are massless).
HANG_HALF_FACTOR = {"abdomen": 0.75, "pelvis": 0.6, "thorax": 0.5, "hand_right": 0.5, "hand_left": 0.5}


def derive_keys(values: dict) -> dict:
    """{name: float} -> {derived name: float}. A key is only produced when its inputs exist.
      <seg>_<s>_radius, _cap_top_z, _cap_bot_z   thigh/shank/upper_arm/forearm capsule: tip at the
                                                 proximal joint (z=0), tip at the distal joint (z=-length)
      knee_<s>_sphere_radius                     sphere centred on the knee (shank frame origin)
      <abdomen|pelvis|thorax|hand_<s>>_mid_z, _half_z   ellipsoid covering the segment (see HANG_HALF_FACTOR)
      thorax_chest_z, _chest_half_z              upper-chest ellipsoid that fills out the shoulders
      shoulder_radius, trapezius_radius, trapezius_start_z   shoulder girdle and the slope to the neck
      foot_<s>_box_pos_x / _half_x               collision box from L/4 behind the ankle to L ahead (centre 3L/8, half 5L/8)
      foot_<s>_box_pos_z / _half_z               sole on the floor: box from the ankle height down to z=0
      foot_<s>_sole_z                            -(ankle height): the sole, for site_sole_<s>
      pelvis_radius                              hip girdle capsule
      head_neck_head_z / _head_half_z            head ellipsoid whose top is the vertex
      head_neck_neck_radius / _start_z           neck capsule from the cervical joint into the head"""
    out = {}
    for side in ("right", "left"):
        for seg, r in CAPSULE_RADII_M.items():
            if (L := values.get(f"{seg}_{side}_length")) is not None:
                out[f"{seg}_{side}_radius"] = r
                out[f"{seg}_{side}_cap_top_z"] = -r
                out[f"{seg}_{side}_cap_bot_z"] = -(L - r)
        if f"shank_{side}_length" in values:
            out[f"knee_{side}_sphere_radius"] = KNEE_SPHERE_RADIUS_M
        if (L := values.get(f"foot_{side}_length")) is not None:
            rear = FOOT_REAR_FRACTION * L  # behind the ankle
            out[f"foot_{side}_box_pos_x"] = (L - rear) / 2  # box from -rear to +L (ankle to the metatarsal heads)
            out[f"foot_{side}_box_half_x"] = (L + rear) / 2
        chain = [values.get(k) for k in ("pelvis_pos_z", f"thigh_{side}_pos_z", f"shank_{side}_pos_z", f"foot_{side}_pos_z")]
        if None not in chain:
            ankle = sum(chain)  # ankle height above the floor in the neutral pose
            out[f"foot_{side}_box_pos_z"] = -ankle / 2
            out[f"foot_{side}_box_half_z"] = ankle / 2
            out[f"foot_{side}_sole_z"] = -ankle  # site_sole, in the foot frame
    for seg, factor in HANG_HALF_FACTOR.items():
        if (L := values.get(f"{seg}_length")) is not None:
            out[f"{seg}_mid_z"] = -L / 2
            out[f"{seg}_half_z"] = factor * L
    if "pelvis_length" in values:
        out["pelvis_radius"] = PELVIS_RADIUS_M
    if (L := values.get("thorax_length")) is not None:
        out["thorax_chest_z"] = -CHEST_Z_FRACTION * L
        out["thorax_chest_half_z"] = CHEST_HALF_FRACTION * L
        out["shoulder_radius"] = SHOULDER_RADIUS_M
        out["trapezius_radius"] = TRAPEZIUS_RADIUS_M
        out["trapezius_start_z"] = -TRAPEZIUS_RADIUS_M
    if (L := values.get("head_neck_length")) is not None:
        out["head_neck_head_z"] = L - HEAD_HALF_HEIGHT_M
        out["head_neck_head_half_z"] = HEAD_HALF_HEIGHT_M
        out["head_neck_neck_radius"] = NECK_RADIUS_M
        out["head_neck_neck_start_z"] = NECK_RADIUS_M
    return out


def warn_unused(text: str, values: dict, report: Report):
    """Flag sheet data the template ignores, so a new sheet value can't silently
    go nowhere. (Unused BSIP values like 'length' are normal; position and ROM
    values are not, unless a position is zero.)"""
    used = {p.key for p in find_placeholders(text)}
    for k, p in sorted(values.items()):
        if k in used or p.table == "bsip":
            continue
        if p.table == "pos" and abs(p.value) < 1e-9:
            continue
        report.warn(f"sheet value {k} = {fmt(p.value)} ({p.source}) is not used by the template")


def add_banner(xml: str, snapshot_sha256: str | None = None) -> str:
    inputs = f"models/{TEMPLATE.name} and the Google Sheet"
    if snapshot_sha256:  # CONVENTIONS section 4: the banner carries the snapshot hash
        inputs += f" (snapshot sha256:{snapshot_sha256[:12]})"
    banner = (f"<!-- GENERATED by scripts/pre_processor.py from {inputs}. "
              "Do not edit; change the inputs and rebuild. -->")
    decl = re.match(r"\s*<\?xml[^>]*\?>\s*", xml)
    if decl:  # a comment may not come before the XML declaration
        return xml[:decl.end()] + banner + "\n" + xml[decl.end():]
    return banner + "\n" + xml


# =============================================================================
# 4. CHECKING THE RESULT
# =============================================================================


def check_xml(xml: str, report: Report):
    """Well-formed? Do the joint ranges make sense?"""
    try:
        root = ET.fromstring(xml)
    except ET.ParseError as e:
        line = xml.splitlines()[e.position[0] - 1].strip() if e.position[0] <= len(xml.splitlines()) else ""
        report.error(f"generated XML is not well-formed: {e}\n          -> {line}")
        return
    for j in root.iter("joint"):
        rng = j.get("range")
        if rng is None:
            continue
        try:
            lo, hi = (float(t) for t in rng.split())
        except ValueError:
            report.error(f"joint '{j.get('name')}': range='{rng}' is not two numbers")
            continue
        if lo >= hi:
            report.error(f"joint '{j.get('name')}': range min {lo:g} is not below max {hi:g}")
        elif lo > 0 or hi < 0:
            report.warn(f"joint '{j.get('name')}': range {lo:g}..{hi:g} excludes the zero pose")


def compile_with_mujoco(xml, report: Report):
    """The real test: can MuJoCo load it? `xml` is the text, or a Path to a file
    (preferred: it tests the exact bytes that will be installed). Returns the
    model or None."""
    try:
        import mujoco
    except ImportError:
        report.warn("mujoco is not installed here, so the compile check was skipped")
        return None
    try:
        if isinstance(xml, Path):
            return mujoco.MjModel.from_xml_path(str(xml))
        return mujoco.MjModel.from_xml_string(xml)
    except Exception as e:  # mujoco raises plain ValueError with its own message
        report.error(f"MuJoCo could not compile the model: {e}")
        return None


# =============================================================================
# 5. ARCHIVE
# =============================================================================


def values_csv(values: dict) -> str:
    """placeholder,value,sheet_source sorted by placeholder: the archive and snapshot format."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["placeholder", "value", "sheet_source"])
    for k in sorted(values):
        w.writerow([k, fmt(values[k].value), values[k].source])
    return buf.getvalue()


def sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sheet_fingerprint(tables: dict) -> str:
    """A short hash of the cell contents of the named ranges."""
    blob = json.dumps({name: t.rows for name, t in sorted(tables.items())}, default=str)
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()[:12]


def pre_processor_commit() -> str:
    """The commit that last changed this script ('unknown' outside git). Stable
    across unrelated commits, so it does not churn the snapshot metadata."""
    try:
        out = subprocess.run(["git", "log", "-1", "--format=%H", "--", Path(__file__).name], cwd=HERE,
                             capture_output=True, text=True, timeout=10, check=True).stdout.strip()
    except (OSError, subprocess.SubprocessError):
        return "unknown"
    return out or "unknown"


def snapshot_meta(snapshot_bytes: bytes, tables: dict, template_text: str, model_bytes: bytes,
                  commit: str) -> str:
    """snapshot.meta.json content. No sheet ID, credential or absolute path goes in here."""
    meta = {"schema": SNAPSHOT_SCHEMA,
            "snapshot_sha256": sha256_hex(snapshot_bytes),
            "sheet_fingerprint": sheet_fingerprint(tables),
            "template_sha256": sha256_hex(template_text.encode("utf-8")),  # text, so line endings can't matter
            "model_sha256": sha256_hex(model_bytes),
            "pre_processor_commit": commit}
    return json.dumps(meta, indent=2, sort_keys=True) + "\n"


def archive_values(values: dict, directory: Path) -> str:
    """Write every value used to a timestamped CSV, but only if it differs from
    the newest one, so the archive is a history of real changes."""
    text = values_csv(values)

    directory.mkdir(exist_ok=True)
    existing = sorted(directory.glob("freyja_params_*.csv"))
    if existing and existing[-1].read_text(encoding="utf-8") == text:
        return f"archive not updated: values identical to {existing[-1].name}"
    path = directory / f"freyja_params_{datetime.now():%Y%m%d_%H%M%S}.csv"
    path.write_text(text, encoding="utf-8")
    shown = path.relative_to(ROOT) if ROOT in path.parents else path
    return f"archived values to {shown}"


# =============================================================================
# 6. MAIN
# =============================================================================


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="Build freyja.xml from the template + Google Sheet.")
    ap.add_argument("--xlsx", metavar="FILE", help="read a downloaded copy of the sheet instead of Google")
    ap.add_argument("--strict", action="store_true", help="missing ROM values are errors, not unlimited joints")
    ap.add_argument("--dry-run", action="store_true", help="check everything but write no files")
    ap.add_argument("--no-archive", action="store_true", help="skip the CSV archive")
    ap.add_argument("--no-checks", action="store_true", help="skip the model checks (gate and advisory)")
    ap.add_argument("--template", type=Path, default=TEMPLATE)
    ap.add_argument("--output", type=Path, default=OUTPUT)
    return ap.parse_args(argv)


def build_values(tables: dict, report: Report):
    """All parsing + data checks. Returns (values, blank_rom_names, ankle_height)."""
    bsip = read_segment_table(tables[RANGE_BSIP], BSIP_COLUMNS, report)
    pos = read_segment_table(tables[RANGE_POS], POS_COLUMNS, report)
    rom, rom_blank = read_rom_table(tables[RANGE_ROM], report)
    check_bsip(bsip, report)
    report.stop_if_errors("reading the sheet")

    values = {}
    for part in (flatten(mirror(bsip, FLIP_BSIP)), flatten(mirror(pos, FLIP_POS)), rom):
        clash = values.keys() & part.keys()
        if clash:
            report.error(f"the same name comes from two tables: {sorted(clash)}")
        values.update(part)
    ankle = check_positions(values, report)
    report.stop_if_errors("checking the sheet values")
    return values, rom_blank, ankle


@dataclass
class BuildResult:
    status: str  # created | updated | unchanged | checked (dry run)
    placeholders: int  # how many placeholders were filled in
    warnings: list
    model_line: str | None  # e.g. "16 bodies, 32 joints, total mass 64.935 kg"
    ankle: float | None
    notes: list
    checks: object = None  # checklib.RunReport when a post_checks hook ran


def build_model(tables: dict, *, template: Path = TEMPLATE, output: Path = OUTPUT,
                failed_output: Path = FAILED_OUTPUT, archive_dir: Path = ARCHIVE_DIR,
                snapshot_dir: Path | None = None, strict: bool = False, dry_run: bool = False,
                archive: bool = True, progress=None, post_checks=None) -> BuildResult:
    """Sheet tables + template -> checked freyja.xml. This is the one function
    both the command line and the watcher use.

    The new model is built next to the real one as a temporary file, checked
    (data, XML, MuJoCo compile), and only then swapped in with os.replace, so
    `output` is never half-written and a bad build never replaces a good one.
    With `snapshot_dir`, snapshot.csv and snapshot.meta.json are staged and
    installed together with the model, or not at all. None = no snapshot (tests,
    custom --output), so a scratch build never overwrites the tracked snapshot.
    `post_checks(xml_path, snapshot_text) -> RunReport` (see checks/checklib.make_post_checks)
    is called on the finished temporary XML before anything is installed. A gate failure parks
    the new XML as `failed_output`, leaves the last good model and snapshot in place and raises
    BuildError naming the failing check IDs; advisory results are returned in BuildResult.checks.
    Raises BuildError (with every problem found) instead of exiting."""
    say = progress or (lambda msg: None)
    report = Report()

    values, rom_blank, ankle = build_values(tables, report)
    say(f"      {sum(k.endswith('_mass') for k in values)} segments (left/right counted separately), "
        f"{len(values)} values, data checks passed")

    say(f"[2/4] Filling template {template.name} ...")
    if not template.exists():
        raise BuildError("reading the template", [f"Template not found: {template}"], report.warnings)
    text = template.read_text(encoding="utf-8")
    plain = {k: p.value for k, p in values.items()}
    plain.update(derive_keys(plain))
    xml, filled = render_counted(text, plain, report, strict=strict)
    report.stop_if_errors("filling the template")
    warn_unused(text, values, report)
    snapshot_bytes = values_csv({**values, **read_targets(tables.get(RANGE_TARGETS), report)}).encode("utf-8")
    xml = add_banner(xml, sha256_hex(snapshot_bytes))

    say("[3/4] Checking the finished XML ...")
    check_xml(xml, report)
    tmp = output.with_name(f".{output.name}.{os.getpid()}.tmp")
    snap_dest = meta_dest = snap_tmp = meta_tmp = None
    if snapshot_dir is not None:
        snap_dest, meta_dest = snapshot_dir / "snapshot.csv", snapshot_dir / "snapshot.meta.json"
        snap_tmp = snap_dest.with_name(f".{snap_dest.name}.{os.getpid()}.tmp")
        meta_tmp = meta_dest.with_name(f".{meta_dest.name}.{os.getpid()}.tmp")
    model = None
    try:
        if not report.errors:
            if dry_run:
                model = compile_with_mujoco(xml, report)
            else:
                tmp.write_bytes(xml.encode("utf-8"))  # bytes: LF on every OS
                model = compile_with_mujoco(tmp, report)
        if report.errors:
            if not dry_run:
                failed_output.write_text(xml, encoding="utf-8")
            report.stop_if_errors(f"checking the XML (the broken build is saved as "
                                  f"{failed_output.name}; {output.name} was NOT touched)")

        checks = None
        if post_checks is not None:
            say("      Running the model checks ...")
            checks = _run_post_checks(post_checks, tmp, xml, snapshot_bytes, dry_run)
            if checks.blocking:
                if not dry_run:
                    failed_output.write_bytes(xml.encode("utf-8"))
                raise BuildError(f"running the gate checks (the new build is saved as {failed_output.name}; "
                                 f"{output.name} and the snapshot were NOT touched)",
                                 [f"{r.id}: {r.message}" for r in checks.blocking], report.warnings)

        say("[4/4] Writing ...")
        notes = []
        if dry_run:
            status = "checked"
            notes.append("dry run: nothing written")
        else:
            staged = [(tmp, output)]
            if snapshot_dir is not None:
                try:
                    snapshot_dir.mkdir(parents=True, exist_ok=True)
                    snap_tmp.write_bytes(snapshot_bytes)
                    meta_tmp.write_bytes(snapshot_meta(snapshot_bytes, tables, text, tmp.read_bytes(),
                                                       pre_processor_commit()).encode("utf-8"))
                except OSError as e:
                    raise BuildError("staging the snapshot", [f"could not write the snapshot ({e}); "
                                     "nothing was installed"], report.warnings, transient=True)
                staged += [(snap_tmp, snap_dest), (meta_tmp, meta_dest)]
            old = output.read_bytes() if output.exists() else None
            changed = [(t, d) for t, d in staged if not d.exists() or d.read_bytes() != t.read_bytes()]
            if not changed:  # byte-identical: leave the files (and their timestamps) alone
                status = "unchanged"
                notes.append(f"{output.name} already up to date (sheet and template give the same model, nothing to write)")
            else:
                _install_together(changed, report)
                status = "created" if old is None else "updated"
                notes.append(f"{output.name} {status}" if (tmp, output) in changed
                             else f"{output.name} unchanged; snapshot metadata {status}")
            if archive:
                try:
                    notes.append(archive_values(values, archive_dir))
                except OSError as e:  # a full disk must not make a good model look failed
                    report.warn(f"could not write the CSV archive: {e}")
    finally:
        for leftover in (tmp, snap_tmp, meta_tmp):
            if leftover is not None:
                leftover.unlink(missing_ok=True)

    model_line = None
    if model is not None:
        model_line = (f"{model.nbody - 1} bodies, {model.njnt} joints, "
                      f"total mass {sum(model.body_mass):.3f} kg")
    return BuildResult(status, filled, report.warnings, model_line, ankle, notes, checks)


def _run_post_checks(post_checks, tmp: Path, xml: str, snapshot_bytes: bytes, dry_run: bool):
    """Call the hook on the exact bytes that would be installed. A dry run has no temporary file next
    to the model, so the hook gets one in the system temp directory, removed afterwards."""
    snapshot_text = snapshot_bytes.decode("utf-8")
    if not dry_run:
        return post_checks(tmp, snapshot_text)
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "freyja.xml"
        path.write_bytes(xml.encode("utf-8"))
        return post_checks(path, snapshot_text)


def _install_together(changed: list, report: Report):
    """Install staged files (tmp, destination) as one unit: if any replace fails,
    the ones already installed are put back and the error is raised. The model
    goes first because it is the file Windows is most likely to have locked."""
    done = []  # (destination, previous bytes or None)
    try:
        for t, d in changed:
            before = d.read_bytes() if d.exists() else None
            _replace_with_retry(t, d, report)
            done.append((d, before))
    except BuildError:
        for d, before in reversed(done):
            if before is None:
                d.unlink(missing_ok=True)
            else:
                d.write_bytes(before)
        raise


def _replace_with_retry(tmp: Path, output: Path, report: Report, tries: int = 6):
    """os.replace is atomic, but Windows refuses while another program has the
    file open (an editor, a viewer). Retry briefly, then give up politely."""
    for attempt in range(tries):
        try:
            os.replace(tmp, output)
            return
        except PermissionError:
            if attempt < tries - 1:
                time.sleep(0.4)
    raise BuildError(f"installing {output.name}",
                     [f"{output.name} is in use by another program, so it could not be replaced. "
                      "The previous version was left untouched. Close whatever has it open."],
                     report.warnings, transient=True)


def default_post_checks(args):
    """The model checks (checks/checklib.py), for the default output only: a scratch build with
    --output or --no-checks skips them, like the snapshot."""
    if args.no_checks or args.output != OUTPUT:
        return None
    sys.path.insert(0, str(ROOT.parent / "checks"))
    import checklib
    return checklib.make_post_checks()


def print_checks(report, say=print):
    """Advisory and waived results and the waiver warnings; passing checks are just counted."""
    passed = sum(r.status == "PASS" for r in report.results)
    say(f"Checks: {passed} of {len(report.results)} passed")
    for r in report.results:
        if r.status != "PASS":
            say(f"  {r.status:7} {r.tier:8} {r.id}: {r.message}")
    for w in report.warnings:
        say(f"  WARNING {w}")


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")  # never crash on an odd character in the console
    args = parse_args(argv)
    try:
        source = f"xlsx file {Path(args.xlsx).name}" if args.xlsx else "Google Sheet"
        print(f"[1/4] Reading {source} ...")
        tables = fetch_xlsx(args.xlsx) if args.xlsx else fetch_google()
        result = build_model(tables, template=args.template, output=args.output, strict=args.strict,
                             dry_run=args.dry_run, archive=not args.no_archive, progress=print,
                             snapshot_dir=PARAMS_DIR if args.output == OUTPUT else None,
                             post_checks=default_post_checks(args))
    except SheetAccessError as e:
        sys.exit(str(e))
    except BuildError as e:
        print("\n" + e.format())
        sys.exit(1)

    print("\n" + "=" * 70)
    if result.model_line:
        print(f"MuJoCo compiled OK: {result.model_line}")
    if result.ankle is not None:
        print(f"Standing pose: ankle joint {result.ankle * 1000:.1f} mm above the floor")
    for n in result.notes:
        print(n)
    if result.checks is not None:
        print_checks(result.checks)
    if result.warnings:
        print(f"\n{len(result.warnings)} warning(s):")
        for w in result.warnings:
            print(f"  WARNING {w}")
    print("=" * 70)


if __name__ == "__main__":
    main()
